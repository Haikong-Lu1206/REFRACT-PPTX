from __future__ import annotations

import copy
from dataclasses import replace
from xml.etree import ElementTree as ET

import pytest

from refract_pptx.design.compiler import compile_proposal
from refract_pptx.design.proposal import AgentProposal, ProposalError
from refract_pptx.evaluation import evaluate_candidate
from refract_pptx.mutation import apply_mutations
from refract_pptx.mutation.ooxml import PackageEditor
from refract_pptx.presentation import object_inventory
from refract_pptx.presentation.objects import A_NS as A
from refract_pptx.presentation.objects import C_NS as C
from refract_pptx.presentation.visual import visual_similarity
from refract_pptx.validation.redteam import run_redteam


def proposal(
    operation,
    component,
    shape_id=4,
    family="reference_reconstruction",
    capability="picture_transform_repair",
):
    return AgentProposal.from_dict(
        {
            "proposal_version": "1.0",
            "title": "Recover the native visual relationship",
            "instruction": "Restore the damaged objects to match the reference. "
            "Keep the presentation editable and preserve unrelated content.",
            "deck_summary": "A synthetic deck with native objects for visual verification.",
            "preservation_contracts": ["Preserve all unrelated visible content."],
            "mutations": [
                {
                    "mutation_id": "visual-repair",
                    "family": family,
                    "capability": capability,
                    "slide": 1,
                    "target": {"shape_id": shape_id},
                    "operation": operation,
                    "scoring": {component: 1.0},
                    "weight": 1.0,
                    "evidence_tier": "reference_visible",
                    "evidence": [
                        {
                            "source": "reference",
                            "locator": "slide 1 target",
                            "supports": "The visible native appearance and relationship.",
                        }
                    ],
                    "rationale": "This corruption changes a visible, recoverable property.",
                    "accepted_solutions": [
                        "Repair or recreate an equivalent editable native object."
                    ],
                }
            ],
        }
    )


@pytest.fixture
def styled_deck(synthetic_pptx, workspace_tmp):
    editor = PackageEditor(synthetic_pptx)
    root = editor.xml("ppt/slides/slide1.xml")
    for paragraph in root.findall(f".//{{{A}}}p"):
        props = ET.Element(f"{{{A}}}pPr", {"algn": "l", "marL": "254000", "indent": "-127000"})
        ET.SubElement(props, f"{{{A}}}buChar", {"char": "\u2022"})
        paragraph.insert(0, props)
    shape = root.find(f".//{{{A}}}r/..")
    assert shape is not None
    for run in root.findall(f".//{{{A}}}r"):
        props = ET.Element(f"{{{A}}}rPr", {"sz": "2400", "b": "0", "i": "0", "u": "none"})
        run.insert(0, props)
        fill = ET.SubElement(props, f"{{{A}}}solidFill")
        ET.SubElement(fill, f"{{{A}}}srgbClr", {"val": "123456"})
    properties = root.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    line = ET.SubElement(properties, f"{{{A}}}ln", {"w": "25400"})
    ET.SubElement(ET.SubElement(line, f"{{{A}}}solidFill"), f"{{{A}}}srgbClr", {"val": "123456"})
    ET.SubElement(line, f"{{{A}}}prstDash", {"val": "solid"})
    return editor.write(workspace_tmp / "styled.pptx")


@pytest.mark.parametrize(
    ("operation", "component", "target", "family", "capability"),
    [
        (
            {"type": "set_paragraph_alignment", "paragraph_index": 0, "alignment": "r"},
            "paragraph_alignment",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {
                "type": "set_paragraph_bullet",
                "paragraph_index": 0,
                "mode": "character",
                "character": "-",
            },
            "paragraph_bullet",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {
                "type": "set_paragraph_indent",
                "paragraph_index": 0,
                "margin_points": 80,
                "indent_points": 20,
            },
            "paragraph_indent",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {"type": "set_rotation", "degrees": 35},
            "rotation",
            4,
            "reference_reconstruction",
            "picture_transform_repair",
        ),
        (
            {"type": "set_flip", "horizontal": True},
            "flip",
            4,
            "reference_reconstruction",
            "picture_transform_repair",
        ),
        (
            {"type": "set_picture_crop", "left": 0.3},
            "picture_crop",
            4,
            "reference_reconstruction",
            "picture_transform_repair",
        ),
        (
            {"type": "set_shape_preset", "preset": "ellipse"},
            "shape_preset",
            2,
            "reference_reconstruction",
            "visible_shape_style_repair",
        ),
        (
            {"type": "set_line_style", "rgb": "FF0000", "width_points": 6, "dash": "dot"},
            "line_style",
            2,
            "reference_reconstruction",
            "visible_shape_style_repair",
        ),
        (
            {"type": "set_font_size", "points": 48},
            "font_size",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {"type": "set_text_color", "rgb": "FF0000"},
            "text_color",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {"type": "set_text_emphasis", "attribute": "bold", "value": True},
            "text_emphasis",
            2,
            "reference_reconstruction",
            "text_style_repair",
        ),
        (
            {"type": "set_chart_direction", "value": "bar"},
            "chart_direction",
            6,
            "native_chart_repair",
            "chart_layout_repair",
        ),
        (
            {"type": "set_chart_grouping", "value": "stacked"},
            "chart_grouping",
            6,
            "native_chart_repair",
            "chart_style_repair",
        ),
        (
            {"type": "set_chart_legend_position", "value": "b"},
            "chart_legend_position",
            6,
            "native_chart_repair",
            "chart_layout_repair",
        ),
    ],
)
def test_extended_native_contracts(
    styled_deck, workspace_tmp, operation, component, target, family, capability
):
    plan = compile_proposal(
        proposal(operation, component, target, family, capability), object_inventory(styled_deck)
    )
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    assert evaluate_candidate(init, init, plan).score == 0
    assert evaluate_candidate(styled_deck, init, plan).score == 1
    # Recreated IDs must not invalidate the honest Oracle.
    editor = PackageEditor(styled_deck)
    root = editor.xml("ppt/slides/slide1.xml")
    for node in root.iter():
        if node.tag.endswith("}cNvPr") and node.get("id") == str(target):
            node.set("id", "120")
    recreated = editor.write(workspace_tmp / "recreated.pptx")
    assert evaluate_candidate(recreated, init, plan).score == 1
    report = run_redteam(styled_deck, init, plan.to_dict(), work_directory=workspace_tmp)
    assert report["valid"], report


def test_partial_crop_and_wrong_image_cannot_earn_transform_credit(styled_deck, workspace_tmp):
    operation = {"type": "set_picture_crop", "left": 0.3}
    plan = compile_proposal(proposal(operation, "picture_crop"), object_inventory(styled_deck))
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    partial_plan = plan.to_dict()
    partial_plan["mutations"][0]["operation"]["left"] = 0.06
    partial = apply_mutations(styled_deck, partial_plan, workspace_tmp / "partial.pptx")
    assert 0 < evaluate_candidate(partial, init, plan).score < 1
    from refract_pptx.evaluation.progress import evaluate_snapshots

    oracle = object_inventory(styled_deck)
    wrong = replace(
        oracle,
        objects=tuple(
            replace(
                obj,
                media_sha256="wrong",
                media_signature=tuple(255 - v for v in obj.media_signature),
            )
            if obj.shape_id == 4 and obj.slide == 1
            else obj
            for obj in oracle.objects
        ),
    )
    assert evaluate_snapshots(wrong, object_inventory(init), plan).score == 0


def test_rotation_wrap_and_equivalent_xml_boolean(styled_deck, workspace_tmp):
    plan = compile_proposal(
        proposal({"type": "set_rotation", "degrees": 45}, "rotation"), object_inventory(styled_deck)
    )
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    alternate = plan.to_dict()
    alternate["mutations"][0]["operation"] = {"type": "set_rotation", "degrees": 360}
    candidate = apply_mutations(styled_deck, alternate, workspace_tmp / "equivalent.pptx")
    assert evaluate_candidate(candidate, init, plan).score == 1
    assert visual_similarity("rotation", {"rotation": 359.5}, {"rotation": 0.2}) == 1


def test_run_split_is_equivalent_but_wrong_content_is_not(styled_deck, workspace_tmp):
    plan = compile_proposal(
        proposal(
            {"type": "set_font_size", "points": 48}, "font_size", 2, capability="text_style_repair"
        ),
        object_inventory(styled_deck),
    )
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    editor = PackageEditor(styled_deck)
    root = editor.xml("ppt/slides/slide1.xml")
    paragraph = root.find(f".//{{{A}}}p")
    run = paragraph.find(f"{{{A}}}r")
    duplicate = copy.deepcopy(run)
    text = run.find(f"{{{A}}}t")
    duplicate.find(f"{{{A}}}t").text = text.text[5:]
    text.text = text.text[:5]
    paragraph.append(duplicate)
    candidate = editor.write(workspace_tmp / "split.pptx")
    assert evaluate_candidate(candidate, init, plan).score == 1
    text.text = "WRONG"
    wrong = editor.write(workspace_tmp / "wrong.pptx")
    assert evaluate_candidate(wrong, init, plan).score == 0


@pytest.mark.parametrize(
    "operation",
    [
        {"type": "set_rotation", "degrees": float("nan")},
        {"type": "set_flip", "horizontal": "false"},
        {"type": "set_picture_crop", "left": 0.6, "right": 0.6},
    ],
)
def test_invalid_parameters_rejected_before_mutation(styled_deck, operation):
    from refract_pptx.presentation.visual_mutation import VISUAL_OPERATIONS

    with pytest.raises(ProposalError):
        compile_proposal(
            proposal(operation, VISUAL_OPERATIONS[operation["type"]]), object_inventory(styled_deck)
        )


def test_hidden_font_evidence_is_rejected(synthetic_pptx):
    with pytest.raises(ProposalError, match="explicit"):
        compile_proposal(
            proposal(
                {"type": "set_font_size", "points": 48},
                "font_size",
                2,
                capability="text_style_repair",
            ),
            object_inventory(synthetic_pptx),
        )


def test_catalog_covers_every_registered_operation():
    from refract_pptx.design.catalog import OPERATION_ARGUMENTS
    from refract_pptx.design.proposal import ALLOWED_OPERATIONS

    assert set(OPERATION_ARGUMENTS) == set().union(*ALLOWED_OPERATIONS.values())


def test_capability_cli_reports_executable_contracts(capsys):
    import json

    from refract_pptx.cli import main

    assert main(["capabilities"]) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["operations"]["set_picture_crop"]["score_components"] == ["picture_crop"]
    assert output["plan_version"] == "1.1"


def test_font_size_partial_progress_on_real_candidates(styled_deck, workspace_tmp):
    plan = compile_proposal(
        proposal(
            {"type": "set_font_size", "points": 48}, "font_size", 2, capability="text_style_repair"
        ),
        object_inventory(styled_deck),
    )
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    scores = []
    for size in (48, 28, 26, 24):
        candidate_plan = plan.to_dict()
        candidate_plan["mutations"][0]["operation"]["points"] = size
        candidate = apply_mutations(
            styled_deck, candidate_plan, workspace_tmp / f"size-{size}.pptx"
        )
        scores.append(evaluate_candidate(candidate, init, plan).score)
    assert scores[0] == 0 and scores[-1] == 1
    assert all(a < b for a, b in zip(scores, scores[1:], strict=False))


def test_protected_visual_damage_is_graded(styled_deck, workspace_tmp):
    plan = compile_proposal(
        proposal({"type": "set_picture_crop", "left": 0.3}, "picture_crop"),
        object_inventory(styled_deck),
    )
    init = apply_mutations(styled_deck, plan, workspace_tmp / "init.pptx")
    damage = compile_proposal(
        proposal(
            {"type": "set_font_size", "points": 90}, "font_size", 2, capability="text_style_repair"
        ),
        object_inventory(styled_deck),
    )
    candidate = apply_mutations(styled_deck, damage, workspace_tmp / "collateral.pptx")
    score = evaluate_candidate(candidate, init, plan)
    assert 0.8 < score.score < 1
    assert all(score.hard_gates.values())


def test_native_marker_restoration(synthetic_pptx, workspace_tmp):
    editor = PackageEditor(synthetic_pptx)
    root = editor.xml("ppt/charts/chart1.xml")
    plot = root.find(f".//{{{C}}}barChart")
    plot.tag = f"{{{C}}}lineChart"
    plot.remove(plot.find(f"{{{C}}}barDir"))
    source = editor.write(workspace_tmp / "line.pptx")
    plan = compile_proposal(
        proposal(
            {"type": "set_chart_marker", "series_index": 0, "value": "triangle"},
            "chart_markers",
            6,
            "native_chart_repair",
            "chart_style_repair",
        ),
        object_inventory(source),
    )
    init = apply_mutations(source, plan, workspace_tmp / "init.pptx")
    assert evaluate_candidate(init, init, plan).score == 0
    assert evaluate_candidate(source, init, plan).score == 1


def test_real_chart_workbook_updates_and_cache_only_attack(workspace_tmp):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["North", "South", "West"]
    data.add_series("Revenue", [12, 18, 27])
    shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
    )
    source = workspace_tmp / "workbook.pptx"
    deck.save(source)
    snapshot = object_inventory(source)
    assert snapshot.objects[0].chart["workbook_state"] == "consistent"
    plan = compile_proposal(
        proposal(
            {"type": "set_chart_value", "series_index": 0, "point_index": 1, "value": 200},
            "chart_data",
            shape.shape_id,
            "native_chart_repair",
            "chart_data_repair",
        ),
        snapshot,
    )
    init = apply_mutations(source, plan, workspace_tmp / "init.pptx")
    assert object_inventory(init).objects[0].chart["workbook_state"] == "consistent"
    assert evaluate_candidate(init, init, plan).score == 0
    assert evaluate_candidate(source, init, plan).score == 1
    # Restore only the chart XML while leaving its embedded workbook wrong.
    candidate = PackageEditor(init)
    original = PackageEditor(source)
    part = snapshot.objects[0].chart["part"]
    candidate.parts[part] = original.parts[part]
    attack = candidate.write(workspace_tmp / "cache-only.pptx")
    assert object_inventory(attack).objects[0].chart["workbook_state"] == "inconsistent"
    assert evaluate_candidate(attack, init, plan).score == 0
    # A normal high-level load/save still retains a valid editable chart.
    loaded = Presentation(init)
    loaded.save(workspace_tmp / "resaved.pptx")
    assert (
        object_inventory(workspace_tmp / "resaved.pptx").objects[0].chart["workbook_state"]
        == "consistent"
    )
