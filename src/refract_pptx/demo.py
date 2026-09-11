"""Generate original, disposable example inputs; no corpus assets are distributed."""

from __future__ import annotations

import json
from pathlib import Path

from refract_pptx.build import build_task
from refract_pptx.design import AgentProposal
from refract_pptx.evaluation import evaluate_candidate
from refract_pptx.mutation import apply_mutations


def create_demo(output: str | Path) -> dict:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise ValueError('Install the demo extra: pip install "refract-pptx[demo]"') from exc
    root = Path(output).resolve()
    if root.exists():
        raise ValueError(f"demo output already exists; choose a new directory: {root}")
    root.mkdir(parents=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(720), Pt(405)
    pdf = canvas.Canvas(str(root / "reference.pdf"), pagesize=(720, 405))
    targets = []
    for heading, body in (
        ("A clear plan", "Compare the reference. Restore layout. Keep content editable."),
        (
            "A measured result",
            "Repair the title while preserving the other objects.",
        ),
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for text, x, y, width, height, size in (
            (heading, 48, 48, 620, 60, 30),
            (body, 48, 145, 620, 65, 16),
        ):
            shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
            shape.text_frame.margin_left = shape.text_frame.margin_top = 0
            shape.text_frame.margin_right = shape.text_frame.margin_bottom = 0
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.text = text
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = RGBColor(35, 45, 65)
            if size == 30:
                targets.append(shape.shape_id)
            pdf.setFillColorRGB(35 / 255, 45 / 255, 65 / 255)
            pdf.setFont("Helvetica", size)
            pdf.drawString(x, 405 - y - size, text)
        pdf.showPage()
    pdf.save()
    source = root / "source.pptx"
    deck.save(source)
    mutations = []
    for index, target in enumerate(targets, 1):
        mutations.append(
            {
                "mutation_id": f"title-{index}",
                "family": "spatial_structure_repair",
                "capability": "object_alignment",
                "slide": index,
                "target": {"shape_id": target, "kind": "shape"},
                "operation": {"type": "move_shape", "dx_points": 55, "dy_points": 30},
                "evidence_tier": "reference_visible",
                "evidence": [
                    {
                        "source": "reference",
                        "locator": f"page {index} title",
                        "supports": "Title aligns with the left edge of the body.",
                    }
                ],
                "rationale": "A visible alignment error with an intact body as an anchor.",
                "accepted_solutions": ["Move or recreate an equivalent editable title."],
                "scoring": {"geometry": 1.0},
                "weight": 0.5,
            }
        )
    proposal = {
        "proposal_version": "1.0",
        "title": "Restore title alignment",
        "instruction": "Restore the two title positions to match reference.pdf. Preserve all "
        "other visible content and keep the presentation editable. Save init.pptx.",
        "deck_summary": "An original two-page layout example, not a difficulty benchmark.",
        "mutations": mutations,
        "preservation_contracts": ["Preserve body text and slide dimensions."],
    }
    (root / "proposal.json").write_text(json.dumps(proposal, indent=2), encoding="utf-8")
    bundle = root / "bundle"
    build_task(
        source,
        root / "reference.pdf",
        AgentProposal.from_dict(proposal).require_valid(),
        bundle,
        task_id="refract-demo",
        source_uri="synthetic://refract-demo",
        license_name="CC0-1.0",
    )
    plan = json.loads((bundle / "evaluator" / "plan.json").read_text())
    partial_plan = {**plan, "mutations": plan["mutations"][1:]}
    partial = apply_mutations(source, partial_plan, root / "one-title-repaired.pptx")
    scores = {
        name: evaluate_candidate(path, bundle / "init.pptx", plan).to_dict()
        for name, path in (
            ("initial", bundle / "init.pptx"),
            ("one_repair", partial),
            ("oracle", source),
        )
    }
    (root / "scores.json").write_text(json.dumps(scores, indent=2), encoding="utf-8")
    return {
        "output": str(root),
        "scores": {key: value["score"] for key, value in scores.items()},
        "note": "Synthetic smoke test; the PDF is authored, not an Office render. "
        "Real tasks still require agent design and target-editor validation.",
    }
